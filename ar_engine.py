# -*- coding: utf-8 -*-
"""
Motor de generación del Análisis Razonado — primera versión.
Cubre: Resumen EBITDA, Análisis Resultado (bridge 6M y 12M), Ingresos (con driver volumen/precio).
"""
import openpyxl
from decimal import Decimal, ROUND_HALF_UP

def load(path):
    return openpyxl.load_workbook(path, data_only=True)

def c(wb, sheet, ref):
    return wb[sheet][ref].value

def _round(value, decimals):
    q = Decimal('1' if decimals == 0 else '1.' + '0'*decimals)
    return Decimal(repr(value)).quantize(q, rounding=ROUND_HALF_UP)

def _grp(s, sep='.'):
    # insert thousands separator into an unsigned integer string
    return f"{int(s):,}".replace(',', sep)

# ---------- formatting (es-CL) ----------
def fmt_millones(raw_thousands, force_decimals=None):
    v = raw_thousands / 1000.0
    neg = v < 0
    av = abs(v)
    decimals = force_decimals if force_decimals is not None else (2 if av < 1000 else 0)
    q = _round(av, decimals)
    if decimals == 0:
        s = _grp(q)
    else:
        whole, frac = str(q).split('.')
        s = _grp(whole) + ',' + frac
    return ('-' if neg else '') + s

def fmt_pct(frac, decimals=2):
    v = frac * 100
    neg = v < 0
    q = _round(abs(v), decimals)
    s = str(q).replace('.', ',') if decimals else str(q)
    return ('-' if neg else '') + s

def fmt_veces(v, decimals=2):
    neg = v < 0
    q = _round(abs(v), decimals)
    s = str(q).replace('.', ',')
    return ('-' if neg else '') + s

def sign_str(v):
    return '+' if v >= 0 else ''  # negative numbers already carry '-' via fmt_millones

def inc_dec(v, inc='incremento', dec='disminución'):
    return inc if v >= 0 else dec

def crec_caida(v, inc='crecimiento', dec='caída'):
    return inc if v >= 0 else dec

# ============================================================
# SECTION: RESUMEN DEL PERIODO (EBITDA)
# ============================================================
def p_ebitda_6m(wb):
    S = 'EBITDA'
    ebitda = c(wb, S, 'C13'); ebitda_prior = c(wb, S, 'E13'); delta = c(wb, S, 'K13'); pct = c(wb, S, 'L13')
    ing_pct = c(wb, S, 'L4'); cost_pct = c(wb, S, 'L21')
    cr_cur = c(wb, S, 'C22'); cr_prior = c(wb, S, 'E22')
    return (
        f"El EBITDA a Jun26 alcanzó US${fmt_millones(ebitda)} millones, lo que representa un "
        f"{inc_dec(pct,'incremento','disminución')} del {fmt_pct(pct)}% al compararlo con los "
        f"US${fmt_millones(ebitda_prior)} millones registrados a Jun25 ({sign_str(delta)}US${fmt_millones(delta)} millones). "
        f"Dicho {inc_dec(pct,'incremento','disminución')} se explica principalmente por el {crec_caida(ing_pct)} en los "
        f"ingresos del {fmt_pct(ing_pct)}%, parcialmente compensado por el {inc_dec(cost_pct,'incremento','disminución')} en los "
        f"Costos y Gastos (excluyendo deterioro y depreciaciones) del {fmt_pct(cost_pct)}%. "
        f"Los costos y gastos sin depreciación representaron un {fmt_pct(abs(cr_cur))}% de los ingresos de actividades "
        f"ordinarias a Jun26, comparado con un {fmt_pct(abs(cr_prior))}% a Jun25."
    )

def p_ebitda_sinfv_6m(wb):
    S = 'EBITDA'
    ebitda = c(wb, S, 'C14'); prior = c(wb, S, 'E14'); pct = c(wb, S, 'L14')
    fv_cur = c(wb, S, 'C18'); fv_prior = c(wb, S, 'E18')
    return (
        f"El EBITDA acumulado a Jun26 sin efecto de fair value de fruta alcanzó US${fmt_millones(ebitda)} millones, "
        f"lo que representa un {inc_dec(pct)} del {fmt_pct(pct)}%, en comparación a los US${fmt_millones(prior)} millones "
        f"obtenidos en el mismo periodo. El impacto de fair value neto a Jun26 fue de US${fmt_millones(fv_cur)} millones, "
        f"mientras que a Jun25 fue de US${fmt_millones(fv_prior)} millones."
    )

# ============================================================
# SECTION: ANALISIS DEL ESTADO DE RESULTADOS (bridge)
# ============================================================
GA_ROWS = {
    'EBITDA': 4, 'Deterioro': 5, 'Depreciación y amortizaciones': 6,
    'Otros ingresos (gastos) no operacionales': 7, 'Costos Financieros Netos': 8,
    'Participación en asociadas': 9, 'Diferencia de cambio': 10,
    'Gasto por impuestos a las ganancias': 11,
}

def ga_row(wb, label, period='6m'):
    """period: '6m' -> cols C(cur)/E(prior)/K(delta)  |  '12m' -> cols G(cur)/I(prior)/N(delta)"""
    r = GA_ROWS[label]
    S = 'Ganancia Atribuible'
    if period == '6m':
        return dict(cur=c(wb,S,f'C{r}'), prior=c(wb,S,f'E{r}'), delta=c(wb,S,f'K{r}'))
    else:
        return dict(cur=c(wb,S,f'G{r}'), prior=c(wb,S,f'I{r}'), delta=c(wb,S,f'N{r}'))

def p_ganancia_controladora_6m(wb):
    d = ga_row(wb, None, '6m') if False else None
    S='Ganancia Atribuible'
    cur=c(wb,S,'C13'); prior=c(wb,S,'E13'); delta=c(wb,S,'K13')
    mejora_empeora = 'una mejora de' if delta>=0 else 'un empeoramiento de'
    ganancia_perdida = lambda v: 'una ganancia' if v>=0 else 'una pérdida'
    return (
        f"Durante los 6 meses terminados en Jun26 se registró {ganancia_perdida(cur)} atribuible a los propietarios "
        f"de la controladora de US${fmt_millones(cur)} millones, que se compara con {ganancia_perdida(prior)} "
        f"controladora registrada a Jun25 por US${fmt_millones(prior)} millones (representando {mejora_empeora} "
        f"{sign_str(delta)}US${fmt_millones(delta)} millones)."
    )

def p_ebitda_bullet_6m(wb):
    """Viñeta de EBITDA del bridge 6M (siempre primero en la lista, no viene de GA_ROWS)."""
    S = 'EBITDA'
    cur = c(wb, S, 'C13'); prior = c(wb, S, 'E13'); delta = c(wb, S, 'K13')
    vol_g = c(wb, 'Ingresos', 'E33'); price_g = c(wb, 'Ingresos', 'D36')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'un mayor volumen comercializado y mayores precios',
        'vol_only': 'un mayor volumen comercializado',
        'vol_despite_price': 'un mayor volumen comercializado',
        'price_only': 'mayores precios',
        'price_despite_vol': 'mayores precios',
        'both_down': 'un menor volumen comercializado y menores precios',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    mayor_menor = 'Mayor' if delta >= 0 else 'Menor'
    return (
        f"{mayor_menor} EBITDA que ascendió a US${fmt_millones(cur)} millones a Jun26, en comparación con "
        f"US${fmt_millones(prior)} millones a Jun25 ({sign_str(delta)}US${fmt_millones(delta)} millones), "
        f"explicado por {phrase}."
    )

def p_ebitda_bullet_12m(wb):
    """Viñeta de EBITDA del bridge 12M (siempre primero en la lista, no viene de GA_ROWS)."""
    S = 'EBITDA'
    cur = c(wb, S, 'G13'); prior = c(wb, S, 'I13'); delta = c(wb, S, 'N13')
    return (
        f"EBITDA de los 12M T25/26 por US${fmt_millones(cur)} millones, con un "
        f"{inc_dec(delta,'incremento','disminución')} con respecto al resultado de los 12M T24/25 por "
        f"US${fmt_millones(prior)} millones ({sign_str(delta)}US${fmt_millones(delta)} millones) por las "
        f"razones ya explicadas."
    )

# Frases de entrada por factor del bridge de Ganancia Atribuible, conscientes de la dirección
# (si el factor mejoró o empeoró el resultado) para que sigan siendo correctas en trimestres futuros.
def bridge_factor_lead_in(label, cur, prior, delta):
    if label == 'Deterioro':
        return 'Menor deterioro del valor de activos que se registró por' if delta >= 0 else \
               'Mayor deterioro del valor de activos que se registró por'
    if label == 'Depreciación y amortizaciones':
        return 'Mayor gasto por depreciación y amortización que se registró por' if delta < 0 else \
               'Menor gasto por depreciación y amortización que se registró por'
    if label == 'Costos Financieros Netos':
        return 'Mayores gastos financieros netos que ascendieron a' if delta < 0 else \
               'Menores gastos financieros netos que ascendieron a'
    if label == 'Diferencia de cambio':
        return 'Diferencia de cambio negativa por' if cur < 0 else 'Diferencia de cambio positiva por'
    if label == 'Gasto por impuestos a las ganancias':
        return 'Impuestos positivos por' if cur >= 0 else 'Impuestos negativos por'
    if label == 'Otros ingresos (gastos) no operacionales':
        return 'Variación favorable en otros ingresos (gastos) no operacionales por' if delta >= 0 else \
               'Variación desfavorable en otros ingresos (gastos) no operacionales por'
    if label == 'Participación en asociadas':
        return 'Variación favorable en la participación en las ganancias de asociadas por' if delta >= 0 else \
               'Variación desfavorable en la participación en las ganancias de asociadas por'
    return f"{label}: variación por"

def bridge_factor_text(label, cur, prior, delta, period='6m'):
    lead = bridge_factor_lead_in(label, cur, prior, delta)
    cur_lbl = 'a Jun26' if period == '6m' else 'en los 12M T25/26'
    prior_lbl = 'a Jun25' if period == '6m' else 'en los 12M T24/25'
    return (
        f"{lead} US${fmt_millones(cur)} millones {cur_lbl}, en comparación con "
        f"US${fmt_millones(prior)} millones {prior_lbl} ({sign_str(delta)}US${fmt_millones(delta)} millones)"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: completar motivo de negocio si corresponde]]."
    )

def p_ganancia_controladora_12m(wb):
    S = 'Ganancia Atribuible'
    cur = c(wb, S, 'G13'); prior = c(wb, S, 'I13'); delta = c(wb, S, 'N13')
    mejora_empeora = 'mejorando' if delta >= 0 else 'empeorando'
    ganancia_perdida = lambda v: 'una ganancia' if v >= 0 else 'una pérdida'
    return (
        f"La ganancia (pérdida) atribuible a los propietarios de la controladora fue de "
        f"US${fmt_millones(cur)} millones en los 12M T25/26, {mejora_empeora} con respecto a "
        f"US${fmt_millones(prior)} millones en los 12M T24/25 ({sign_str(delta)}US${fmt_millones(delta)} millones), "
        f"cuyo resultado se explica principalmente por:"
    )

def p_ebitda_12m_intro(wb):
    S = 'EBITDA'
    cur = c(wb, S, 'G13'); prior = c(wb, S, 'I13')
    return (
        f"Durante los 12M T25/26, el EBITDA ascendió a US${fmt_millones(cur)} millones, representando un "
        f"incremento con respecto al EBITDA de los 12M T24/25 por US${fmt_millones(prior)} millones. Este "
        f"mayor EBITDA se explica por:"
    )

def p_ebitda_12m_bullet_ingresos(wb):
    pct = c(wb, 'Ingresos', 'M4')
    vol_g = c(wb, 'Ingresos', 'E54'); price_g = c(wb, 'Ingresos', 'E57')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'asociado a mayores volúmenes y precios',
        'vol_only': 'asociado a mayores volúmenes',
        'vol_despite_price': 'asociado a mayores volúmenes',
        'price_only': 'asociado a mayores precios',
        'price_despite_vol': 'asociado a mayores precios',
        'both_down': 'asociado a menores volúmenes y precios',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    return f"{inc_dec(pct,'Incremento','Disminución')} en los ingresos del {fmt_pct(pct)}% {phrase}."

def p_ebitda_12m_bullet_costos(wb):
    S = 'EBITDA'
    cur_ratio = c(wb, S, 'G22'); prior_ratio = c(wb, S, 'I22')
    cost_prior = c(wb, S, 'I21'); cost_delta = c(wb, S, 'N21')
    pct = abs(cost_delta) / abs(cost_prior)
    return (
        f"Por su parte, los costos y gastos sin incluir depreciación registraron un {inc_dec(pct,'aumento','disminución')} "
        f"de {fmt_pct(pct)}%. El total de costos y gastos sin incluir depreciación representaron un "
        f"{fmt_pct(abs(cur_ratio))}% del total de los ingresos en los 12M T25/26, comparado con el "
        f"{fmt_pct(abs(prior_ratio))}% en los 12M T24/25."
    )

def bridge_bullets(wb, period='6m', top_n=None, threshold_abs_mus=3.0):
    """Returns ranked list of (label, cur, prior, delta) by |delta|, excluding items below threshold (in MUS millones)."""
    items = []
    for label in GA_ROWS:
        if label == 'EBITDA':
            continue  # EBITDA always called out separately / first
        d = ga_row(wb, label, period)
        items.append((label, d['cur'], d['prior'], d['delta']))
    items.sort(key=lambda x: abs(x[3]), reverse=True)
    if top_n:
        items = items[:top_n]
    else:
        items = [it for it in items if abs(it[3])/1000.0 >= threshold_abs_mus]
    return items

# ============================================================
# SECTION: INGRESOS (con driver volumen/precio)
# ============================================================
def volumen_precio_driver(vol_growth, price_growth, min_material=0.03):
    """Clasifica el driver de una variación de ingresos según volumen/precio.
    Devuelve una de: 'both', 'vol_only', 'price_only', 'vol_despite_price', 'price_despite_vol', 'both_down', 'unclear'
    """
    vol_up = vol_growth > min_material
    price_up = price_growth > min_material
    vol_down = vol_growth < -min_material
    price_down = price_growth < -min_material
    if vol_up and price_up:
        return 'both'
    if vol_up and price_down:
        return 'vol_despite_price'
    if price_up and vol_down:
        return 'price_despite_vol'
    if vol_up:
        return 'vol_only'
    if price_up:
        return 'price_only'
    if vol_down and price_down:
        return 'both_down'
    return 'unclear'

def p_ingresos_totales_6m(wb):
    S='Ingresos'
    cur=c(wb,S,'C4'); prior=c(wb,S,'E4'); pct=c(wb,S,'G4')
    vol_g = c(wb,S,'E33'); price_g = c(wb,S,'D36')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'debido al incremento en volúmenes y precios de venta',
        'vol_only': 'debido al incremento en los volúmenes de venta',
        'vol_despite_price': 'debido al incremento en los volúmenes de venta',
        'price_only': 'debido al incremento en los precios de venta',
        'price_despite_vol': 'debido al incremento en los precios de venta',
        'both_down': 'debido a la caída en volúmenes y precios de venta',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    return (
        f"Los ingresos de actividades ordinarias alcanzaron US${fmt_millones(cur, force_decimals=0)} millones a Jun26, "
        f"representando un {inc_dec(pct)} del {fmt_pct(pct)}% con respecto a Jun25, {phrase}."
    )

def p_frutafresca_6m(wb):
    S='Ingresos'
    pct = c(wb,S,'G15')
    vol_g = c(wb,S,'E39'); price_g = c(wb,S,'E42')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'explicado por los mayores volúmenes y precios comercializados',
        'vol_only': 'explicado principalmente por los mayores volúmenes comercializados',
        'vol_despite_price': 'explicado principalmente por los mayores volúmenes comercializados',
        'price_only': 'explicado principalmente por los mayores precios',
        'price_despite_vol': 'explicado principalmente por los mayores precios',
        'both_down': 'explicado por los menores volúmenes y precios comercializados',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    return (
        f"Las ventas del segmento de Fruta Fresca a Jun26 {inc_dec(pct,'incrementaron','disminuyeron')} en {fmt_pct(pct)}% "
        f"respecto del período anterior, {phrase}."
    )

def p_valoragregado_6m(wb):
    S='Ingresos'
    pct = c(wb,S,'G22')
    vol_g = c(wb,S,'E45'); price_g = c(wb,S,'E48')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'explicado por aumento de volúmenes y precios',
        'vol_only': 'explicado principalmente por mayores volúmenes',
        'vol_despite_price': 'explicado principalmente por mayores volúmenes',
        'price_only': 'explicado principalmente por mayores precios',
        'price_despite_vol': 'explicado principalmente por mayores precios',
        'both_down': 'explicado por menores volúmenes y precios',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    share_cur = c(wb,S,'C27'); share_prior = c(wb,S,'E27')
    return (
        f"Los productos con valor agregado registraron un incremento en los ingresos por venta a Jun26 del {fmt_pct(pct)}% "
        f"respecto de los ingresos registrados a Jun25, {phrase}. "
        f"Los ingresos del segmento de valor agregado representaron un {fmt_pct(share_cur)}% de los ingresos a Jun26 "
        f"({fmt_pct(share_prior)}% a Jun25)."
    )

def p_ingresos_totales_12m(wb):
    S='Ingresos'
    cur=c(wb,S,'I4'); prior=c(wb,S,'K4'); pct=c(wb,S,'M4')
    vol_g = c(wb,S,'E54'); price_g = c(wb,S,'E57')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'debido a un incremento en los volúmenes y precios de venta',
        'vol_only': 'debido a un incremento en los volúmenes de venta',
        'vol_despite_price': 'debido a un incremento en los volúmenes de venta',
        'price_only': 'debido a un incremento en los precios de venta',
        'price_despite_vol': 'debido a un incremento en los precios de venta',
        'both_down': 'debido a una caída en los volúmenes y precios de venta',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    return (
        f"Los ingresos de actividades ordinarias alcanzaron US${fmt_millones(cur, force_decimals=0)} millones en los "
        f"12M T25/26, representando un {inc_dec(pct)} del {fmt_pct(pct)}% con respecto a los "
        f"US${fmt_millones(prior, force_decimals=0)} millones de los 12M T24/25, {phrase}."
    )

def p_frutafresca_12m(wb):
    S='Ingresos'
    pct = c(wb,S,'M15')
    vol_g = c(wb,S,'E60'); price_g = c(wb,S,'E63')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'explicado por los mayores volúmenes y precios comercializados',
        'vol_only': 'explicado principalmente por mayores volúmenes comercializados',
        'vol_despite_price': 'explicado principalmente por mayores volúmenes comercializados',
        'price_only': 'explicado principalmente por mayores precios',
        'price_despite_vol': 'explicado principalmente por mayores precios',
        'both_down': 'explicado por menores volúmenes y precios comercializados',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    return (
        f"Las ventas del segmento de Fruta Fresca en los 12M T25/26 {inc_dec(pct,'incrementaron','disminuyeron')} en "
        f"{fmt_pct(pct)}% respecto de los 12M T24/25, {phrase}."
    )

def p_valoragregado_12m(wb):
    S='Ingresos'
    pct = c(wb,S,'M22')
    vol_g = c(wb,S,'E67'); price_g = c(wb,S,'E70')
    driver = volumen_precio_driver(vol_g, price_g)
    phrase = {
        'both': 'explicado tanto por mayores volúmenes en este segmento como por mayores precios',
        'vol_only': 'explicado principalmente por mayores volúmenes en este segmento',
        'vol_despite_price': 'explicado principalmente por mayores volúmenes en este segmento',
        'price_only': 'explicado principalmente por mayores precios',
        'price_despite_vol': 'explicado principalmente por mayores precios',
        'both_down': 'explicado por menores volúmenes y precios en este segmento',
        'unclear': '[REVISAR: driver volumen/precio no concluyente]',
    }[driver]
    share_cur = c(wb,S,'I27'); share_prior = c(wb,S,'K27')
    return (
        f"Por su parte, el segmento de productos con valor agregado registró un incremento en los ingresos del "
        f"{fmt_pct(pct)}% durante los 12M T25/26, en comparación con los 12M T24/25, {phrase}. Este segmento "
        f"representó el {fmt_pct(share_cur)}% de los ingresos en los 12M T25/26, en comparación con el "
        f"{fmt_pct(share_prior)}% que representó en los 12M T24/25."
    )

# ============================================================
# SECTION: COSTOS Y GASTOS
# ============================================================
def p_costoventa_6m(wb):
    S='Costos'
    cur=abs(c(wb,S,'D6')); prior=abs(c(wb,S,'F6')); pct=c(wb,S,'H6')
    ratio_cur=c(wb,S,'D27'); ratio_prior=c(wb,S,'F27')
    vol_g = c(wb,S,'H35')  # volumen distribuido var%
    cost_per_kg_cur=c(wb,S,'D36'); cost_per_kg_prior=c(wb,S,'F36')
    # si el costo/kg bajó o se mantuvo, el aumento es por volumen; si el costo/kg también subió, mencionar ambos
    cpk_up = cost_per_kg_cur < cost_per_kg_prior * 1.0 and abs(cost_per_kg_cur) > abs(cost_per_kg_prior) * 1.03
    driver = "asociado al crecimiento en volúmenes y al mayor costo unitario" if cpk_up else "asociado al crecimiento en volúmenes"
    return (
        f"Los costos de ventas a Jun26 totalizaron US${fmt_millones(cur, force_decimals=0)} millones, presentando un "
        f"{inc_dec(pct)} del {fmt_pct(pct)}% respecto a los US${fmt_millones(prior, force_decimals=0)} millones "
        f"registrados a Jun25, {driver}. Los costos de ventas representaron un {fmt_pct(abs(ratio_cur))}% de los "
        f"ingresos de actividades ordinarias a Jun26, mientras que a Jun25 alcanzaban un {fmt_pct(abs(ratio_prior))}%."
    )

def p_costoventa_12m(wb):
    S='Costos'
    cur=abs(c(wb,S,'J6')); prior=abs(c(wb,S,'L6')); pct=c(wb,S,'N6')
    ratio_cur=c(wb,S,'J27'); ratio_prior=c(wb,S,'L27')
    return (
        f"Los costos de ventas de los 12M T25/26 alcanzaron US${fmt_millones(cur, force_decimals=0)} millones, "
        f"representando un {inc_dec(pct)} del {fmt_pct(pct)}% respecto a los US${fmt_millones(prior, force_decimals=0)} "
        f"millones registrados en los 12M T24/25, explicado también por los mayores volúmenes comercializados. "
        f"Los costos de ventas representaron un {fmt_pct(abs(ratio_cur))}% de los ingresos de actividades ordinarias "
        f"en los 12M T25/26, comparado con el {fmt_pct(abs(ratio_prior))}% en los 12M T24/25"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: completar motivo de la variación residual, ej. 'productores terceros']]."
    )

def p_gastosadmin_6m(wb):
    S='Costos'
    cur=abs(c(wb,S,'D8')); prior=abs(c(wb,S,'F8')); pct=c(wb,S,'H8')
    return (
        f"Los gastos de administración a Jun26 alcanzaron US${fmt_millones(cur)} millones, representando un "
        f"{inc_dec(pct)} del {fmt_pct(pct)}% con respecto a Jun25"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: completar motivo, ej. 'provisiones no recurrentes']]."
    )

def p_otrosgastos_6m(wb):
    S='Costos'
    cur=abs(c(wb,S,'D9')); prior=abs(c(wb,S,'F9'))
    return (
        f"Los otros gastos, por función (excluyendo el deterioro del valor de activos) alcanzaron "
        f"US${fmt_millones(cur)} millones a Jun26, en comparación con US${fmt_millones(prior)} millones a Jun25."
    )

def p_deterioro_6m(wb):
    S='Costos'
    cur=abs(c(wb,S,'D13')); prior=abs(c(wb,S,'F13'))
    return (
        f"El gasto por deterioro de valor de activos a Jun26 fue de US${fmt_millones(cur)} millones"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: completar motivo del deterioro Jun26]], en comparación con "
        f"US${fmt_millones(prior)} millones a Jun25"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: completar motivo del deterioro Jun25]]."
    )

def p_otroscomponentes_6m(wb):
    S='Otros ingresos(gastos)'
    cur=abs(c(wb,S,'D13')); prior=abs(c(wb,S,'F13'))
    GA='Ganancia Atribuible'
    cf_cur=c(wb,GA,'C8'); cf_prior=c(wb,GA,'E8'); cf_pct=c(wb,GA,'L8')
    return (
        f"Los otros componentes del resultado registraron un costo de US${fmt_millones(cur)} millones a Jun26, en "
        f"comparación con los US${fmt_millones(prior)} millones a Jun25. La principal partida no operacional son los "
        f"costos financieros netos que ascienden a US${fmt_millones(abs(cf_cur))} millones a Jun26, en comparación "
        f"con US${fmt_millones(abs(cf_prior))} millones a Jun25, registrando un {inc_dec(cf_pct)} del {fmt_pct(cf_pct)}%."
    )

def p_impuesto_6m(wb):
    S='Ganancia Atribuible'
    cur=c(wb,S,'C11'); prior=c(wb,S,'E11')
    return (
        f"A Jun26, se registraron gastos por impuesto a las ganancias por US${fmt_millones(abs(cur))} millones, en "
        f"comparación a los US${fmt_millones(abs(prior))} millones registrados a Jun25."
    )

# ============================================================
# SECTION: INDICADORES DE ACTIVIDAD
# ============================================================
def p_rotacion_activos(wb):
    S='Ind. Actividad'
    cur=c(wb,S,'D5'); prior=c(wb,S,'E5'); ing_pct=c(wb,S,'J3')
    return (
        f"La rotación de los activos al 30 de junio de 2026 {inc_dec(cur-prior,'incrementó','disminuyó')} desde "
        f"{fmt_veces(prior)} veces a Jun25 a {fmt_veces(cur)} veces a Jun26, explicado principalmente por el "
        f"incremento en los ingresos del {fmt_pct(ing_pct)}%."
    )

def p_rotacion_inventarios(wb):
    S='Ind. Actividad'
    cur=c(wb,S,'D7'); prior=c(wb,S,'E7'); inv_prom_pct=c(wb,S,'J8')
    verbo = 'aumentaron' if inv_prom_pct>=0 else 'disminuyeron'
    nivel = 'mayores' if inv_prom_pct>=0 else 'menores'
    return (
        f"Por su parte, el ratio de rotación de inventarios {inc_dec(cur-prior,'incrementó','disminuyó')} desde "
        f"{fmt_veces(prior)} veces a Jun25 a {fmt_veces(cur)} veces a Jun26, explicado por los {nivel} niveles de "
        f"inventario promedio."
    )

# ============================================================
# SECTION: INDICADORES FINANCIEROS Y RENTABILIDAD
# ============================================================
def p_liquidez(wb):
    IF='Ind. Financieros'; BAL='Balance'
    cur=c(wb,IF,'E7'); prior=c(wb,IF,'F7'); pct=c(wb,IF,'G7')
    ac_pct=c(wb,BAL,'H7'); pc_pct=c(wb,BAL,'H11')
    leve = 'leve ' if abs(pct) < 0.02 else ''
    return (
        f"La liquidez corriente fue de {fmt_veces(cur)} veces a Jun26, lo que representa una {leve}"
        f"{inc_dec(pct,'alza','reducción')} de {fmt_pct(abs(pct))}% con respecto a Dic25 ({fmt_veces(prior)} veces), "
        f"explicado por una mayor {inc_dec(ac_pct,'alza','reducción')} de los activos corrientes del período "
        f"({fmt_pct(ac_pct)}%), en relación con la {inc_dec(pc_pct,'alza','reducción')} de los pasivos corrientes "
        f"({fmt_pct(pc_pct)}%)."
    )

def p_razon_acida(wb):
    S='Ind. Financieros'
    cur=c(wb,S,'E9'); prior=c(wb,S,'F9')
    return (
        f"En tanto, la razón ácida alcanzó las {fmt_veces(cur)} veces, experimentando un "
        f"{inc_dec(cur-prior,'incremento','disminución')} en relación con Dic25."
    )

def p_razon_endeudamiento(wb):
    IF='Ind. Financieros'; BAL='Balance'
    cur=c(wb,IF,'E11'); pct=c(wb,IF,'G11')
    pt_pct=c(wb,BAL,'H13')
    return (
        f"La razón de endeudamiento {inc_dec(pct,'aumentó','disminuyó')} en {fmt_pct(abs(pct))}% con respecto a "
        f"Dic25, llegando a {fmt_veces(cur)} veces, explicado principalmente por la "
        f"{inc_dec(pt_pct,'alza','reducción')} de los pasivos totales en un {fmt_pct(abs(pt_pct))}%"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: confirmar si aplica 'asociado principalmente a la estacionalidad del negocio' u otra razón]]."
    )

def p_estructura_deuda(wb):
    S='Ind. Financieros'
    cur_cp=c(wb,S,'E13'); prior_cp=c(wb,S,'F13'); cur_lp=c(wb,S,'E15'); prior_lp=c(wb,S,'F15')
    mantuvo = abs(cur_cp-prior_cp) < 0.03
    verbo = 'mantuvo' if mantuvo else 'modificó'
    return (
        f"Durante este periodo, la compañía {verbo} su estructura de deuda. El porcentaje de pasivos corrientes a "
        f"Jun26 fue de {fmt_pct(cur_cp)}% respecto de los pasivos totales, comparado con el {fmt_pct(prior_cp)}% "
        f"registrado al 31 de diciembre de 2025. La proporción de pasivos de largo plazo pasó de {fmt_pct(prior_lp)}% "
        f"a Dic25 a {fmt_pct(cur_lp)}% a Jun26."
    )

def p_cobertura_gf(wb):
    IR='Ind. Rentabilidad'
    cur=c(wb,IR,'C2'); prior=c(wb,IR,'D2')
    return (
        f"El índice de cobertura de gastos financieros se ubica en {fmt_veces(cur)} veces a Jun26, "
        f"{inc_dec(cur-prior,'mejorando','empeorando')} desde las {fmt_veces(prior)} veces de Jun25, debido a la "
        f"mayor utilidad antes de impuestos. Por su parte, los costos financieros a Jun26 presentan un "
        f"{inc_dec(1,'incremento')} respecto a Jun25"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: confirmar motivo, ej. 'costos de prepagos de deuda realizados durante el período']]."
    )

def p_rentabilidad_patrim_controladora(wb):
    IR='Ind. Rentabilidad'; GA='Ganancia Atribuible'
    cur=c(wb,IR,'C4'); prior=c(wb,IR,'D4')
    gc_cur=c(wb,GA,'C13'); gc_prior=c(wb,GA,'E13')
    return (
        f"La rentabilidad del patrimonio de la controladora {inc_dec(cur-prior,'mejoró','empeoró')} desde "
        f"{fmt_pct(prior)}% a Jun25 a {fmt_pct(cur)}% a Jun26, explicado por el "
        f"{inc_dec(gc_cur-gc_prior,'incremento','deterioro')} en la ganancia controladora, que pasó de "
        f"US${fmt_millones(gc_prior)} millones a Jun25 a US${fmt_millones(gc_cur)} millones a Jun26."
    )

def p_rentabilidad_patrim_total(wb):
    IR='Ind. Rentabilidad'
    cur=c(wb,IR,'C6'); prior=c(wb,IR,'D6')
    ganancia_ejercicio=c(wb,IR,'L10')
    return (
        f"Por su parte, la rentabilidad del patrimonio total a Jun26 se ubica en {fmt_pct(cur)}%, en comparación "
        f"con la rentabilidad de {fmt_pct(prior)}% a Jun25, explicado por la ganancia del ejercicio de "
        f"US${fmt_millones(ganancia_ejercicio)} millones a Jun26."
    )

# ============================================================
# SECTION: DEUDA FINANCIERA NETA
# ============================================================
def p_dfn(wb):
    S='DF Neta'
    cur=c(wb,S,'C15'); prior=c(wb,S,'E15')
    arr_cur=c(wb,S,'H8'); arr_prior=c(wb,S,'I8')
    ifrs16_cur=c(wb,S,'C28'); ifrs16_prior=c(wb,S,'E28')
    leasing_cur=c(wb,S,'C22'); leasing_prior=c(wb,S,'E22')
    return (
        f"La deuda financiera neta de la Sociedad se redujo desde US${fmt_millones(prior)} millones al 31 de "
        f"diciembre de 2025 a US${fmt_millones(cur)} millones al 30 de junio de 2026. Al 30 de junio de 2026, el "
        f"pasivo por arrendamiento asciende a US${fmt_millones(arr_cur)} millones, de los cuales "
        f"US${fmt_millones(ifrs16_cur)} millones corresponden a pasivos por arrendamientos bajo IFRS16 y "
        f"US${fmt_millones(leasing_cur)} millones estaban asociados a pasivos registrados por leasings. Por su "
        f"parte, al 31 de diciembre de 2025, el pasivo por arrendamiento ascendía a US${fmt_millones(arr_prior)} "
        f"millones, de los cuales US${fmt_millones(ifrs16_prior)} millones correspondían a pasivos por "
        f"arrendamientos bajo IFRS16 y US${fmt_millones(leasing_prior)} millones estaban asociados a pasivos por "
        f"leasings."
    )

# ============================================================
# SECTION: ESTADO DE SITUACION FINANCIERA (Balance bridge)
# ============================================================
FECU = 'EEFF Presentación Fecu'

def _fecu_row(wb, label, start=1, end=176):
    ws = wb[FECU]
    for r in range(start, end+1):
        if ws.cell(r, 3).value == label:  # col C
            cur = ws.cell(r,5).value; prior = ws.cell(r,7).value
            delta = ws.cell(r,9).value
            if delta is None and cur is not None and prior is not None:
                delta = cur - prior
            return dict(cur=cur, prior=prior, delta=delta)
    raise KeyError(label)

def p_activos_totales_intro(wb):
    d = _fecu_row(wb, 'Total Activos')
    pct = d['delta']/d['prior']
    return (
        f"Al 30 de junio de 2026, los activos totales se {inc_dec(d['delta'],'incrementaron','redujeron')} en "
        f"US${fmt_millones(abs(d['delta']))} millones ({fmt_pct(pct)}%) con respecto a los existentes al 31 de "
        f"diciembre de 2025, explicado principalmente por:"
    )

ACTIVOS_NC_LABELS = [
    'Otros activos financieros, no corrientes', 'Otros activos no financieros, no corrientes',
    'Derechos por cobrar, no corrientes', 'Cuentas por cobrar a entidades relacionadas, no corrientes',
    'Inversiones contabilizadas utilizando el método de la participación',
    'Activos intangibles distintos de la plusvalía', 'Plusvalía', 'Propiedades, planta y equipo',
    'Activos por Derecho de Uso', 'Activos por impuestos diferidos',
]
ACTIVOS_C_LABELS = [
    'Efectivo y equivalentes al efectivo', 'Otros activos financieros, corrientes',
    'Otros activos no financieros, corrientes', 'Deudores comerciales y otras cuentas por cobrar, corrientes',
    'Cuentas por cobrar a entidades relacionadas, corrientes', 'Inventarios', 'Activos biológicos, corrientes',
    'Activos por impuestos corrientes',
]

def _top_items(wb, labels, top_n=3):
    items = [(lbl, _fecu_row(wb, lbl)) for lbl in labels]
    items.sort(key=lambda x: abs(x[1]['delta']), reverse=True)
    return items[:top_n]

def p_activos_nocorrientes_bridge(wb):
    total = _fecu_row(wb, 'Total Activos no corrientes')
    top = _top_items(wb, ACTIVOS_NC_LABELS, 3)
    ppe_label, ppe = top[0]
    rest_parts = []
    for lbl, d in top[1:]:
        short = {'Plusvalía':'cuentas de Plusvalía', 'Activos intangibles distintos de la plusvalía':'Activos intangibles'}.get(lbl, lbl)
        rest_parts.append(f"reducción de {short} (US${fmt_millones(d['delta'])} millones)" if d['delta']<0 else f"incremento de {short} (US$+{fmt_millones(d['delta'])} millones)")
    rest = ', '.join(rest_parts)
    verbo = 'Reducción' if total['delta'] < 0 else 'Incremento'
    lead = f"{verbo} en activos no corrientes en US${fmt_millones(abs(total['delta']))} millones respecto a Dic25:"
    return (
        "{{" + lead + "}} "
        f"Explicado principalmente por la {inc_dec(ppe['delta'],'alza','reducción')} en {ppe_label} "
        f"(US${fmt_millones(ppe['delta'])} millones) asociado a la depreciación del período, compensado por "
        f"inversiones también del período"
        f" [[CAUSA NO DERIVABLE DEL EXCEL: confirmar frase de contexto si aplica]]. Adicionalmente, se explica por "
        f"{rest}, entre otras variaciones menores."
    )

def p_activos_corrientes_bridge(wb):
    total = _fecu_row(wb, 'Total de activos corrientes distintos de los activos o grupos de activos para su disposición clasificados como mantenidos para la venta')
    top = _top_items(wb, ACTIVOS_C_LABELS, 3)
    negs = [(lbl,d) for lbl,d in top if d['delta']<0]
    poss = [(lbl,d) for lbl,d in top if d['delta']>=0]
    neg_parts = []
    for lbl, d in negs:
        short = {'Cuentas por cobrar a entidades relacionadas, corrientes':'Cuentas por cobrar a entidades relacionadas', 'Inventarios':'nivel de inventarios'}.get(lbl, lbl)
        tag = 'menores' if 'nivel' not in short else 'menor'
        neg_parts.append(f"{tag} {short} (US${fmt_millones(d['delta'])} millones)")
    pos_parts = [f"mayores {lbl} (US$+{fmt_millones(d['delta'])} millones)" for lbl, d in poss]
    verbo = 'Reducción' if total['delta'] < 0 else 'Incremento'
    lead = f"{verbo} en activos corrientes en US${fmt_millones(abs(total['delta']))} millones respecto a Dic25:"
    return (
        "{{" + lead + "}} "
        f"Explicado principalmente por las {' y '.join(neg_parts)}, asociado principalmente a la estacionalidad "
        f"del negocio. Lo anterior, fue compensado parcialmente por los {' y '.join(pos_parts)}."
    )

PASIVOS_GROUPS = {
    'pasivos financieros, corrientes y no corrientes': [
        ('Otros pasivos financieros, corrientes', 1), ('Otros pasivos financieros, no corrientes', 1)],
    'cuentas comerciales por pagar, corrientes y no corrientes': [
        ('Cuentas comerciales y otras cuentas por pagar, corrientes', 1), ('Otras cuentas por pagar, no corrientes', 1)],
    'pasivos por arrendamientos': [
        ('Pasivos por arrendamientos, corrientes', 1), ('Pasivos por arrendamientos, no corrientes', 1)],
}

def p_pasivos_bridge(wb):
    total = _fecu_row(wb, 'Total pasivos')
    group_deltas = {}
    for g, rows in PASIVOS_GROUPS.items():
        tot = 0
        for lbl, _ in rows:
            tot += _fecu_row(wb, lbl)['delta']
        group_deltas[g] = tot
    ordered = sorted(group_deltas.items(), key=lambda x: abs(x[1]), reverse=True)
    main, rest = ordered[0], ordered[1:]
    rest_txt = ', '.join(f"las menores {g} (US${fmt_millones(v)} millones)" if v<0 else f"los mayores {g} (US$+{fmt_millones(v)} millones)" for g,v in rest)
    verbo = 'se redujeron' if total['delta'] < 0 else 'se incrementaron'
    main_verbo = 'reducción' if main[1] < 0 else 'incremento'
    bold_mid = f"pasivos totales {verbo} en US${fmt_millones(abs(total['delta']))} millones"
    return (
        "Los {{" + bold_mid + "}}, explicado "
        f"principalmente por la {main_verbo} en los {main[0]} (US${fmt_millones(main[1])} millones), explicado por la "
        f"estacionalidad del negocio. Adicionalmente, por {rest_txt}, también asociado a la estacionalidad."
    )

def p_patrimonio_bridge(wb):
    patr = _fecu_row(wb, '   Total Patrimonio')
    ganancia_cur = c(wb, FECU, 'E119')  # 'Ganancia (pérdida)' del ejercicio (fila específica; la etiqueta se repite en la hoja)
    bold_mid = (f"patrimonio total de la Compañía {inc_dec(patr['delta'],'incrementó','disminuyó')} en "
                f"US${fmt_millones(abs(patr['delta']))} millones")
    return (
        "El {{" + bold_mid + "}} con respecto al 31 de diciembre de 2025, totalizando "
        f"US${fmt_millones(patr['cur'])} millones, explicado principalmente por la ganancia del período por "
        f"US${fmt_millones(ganancia_cur)} millones."
    )

if __name__ == '__main__':
    import sys
    wb = load(sys.argv[1])
    print("P49:", p_ebitda_6m(wb))
    print()
    print("P51:", p_ebitda_sinfv_6m(wb))
    print()
    print("P58:", p_ganancia_controladora_6m(wb))
    print()
    print("Bridge 6M (top by magnitude, >=3 MUS$):")
    for label, cur, prior, delta in bridge_bullets(wb, '6m'):
        print(f"   {label}: cur={fmt_millones(cur)} prior={fmt_millones(prior)} delta={sign_str(delta)}{fmt_millones(delta)}")
    print()
    print("P97:", p_ingresos_totales_6m(wb))
    print()
    print("P99:", p_frutafresca_6m(wb))
    print()
    print("P101:", p_valoragregado_6m(wb))
    print()
    print("P105:", p_ingresos_totales_12m(wb))
    print()
    print("P107:", p_frutafresca_12m(wb))
    print()
    print("P109:", p_valoragregado_12m(wb))
    print()
    print("P118:", p_costoventa_6m(wb))
    print()
    print("P120:", p_costoventa_12m(wb))
    print()
    print("P124:", p_gastosadmin_6m(wb))
    print()
    print("P128:", p_otrosgastos_6m(wb))
    print()
    print("P132:", p_deterioro_6m(wb))
    print()
    print("P138:", p_otroscomponentes_6m(wb))
    print()
    print("P142:", p_impuesto_6m(wb))
    print()
    print("P148:", p_rotacion_activos(wb))
    print()
    print("P150:", p_rotacion_inventarios(wb))
    print()
    print("P157:", p_dfn(wb))
    print()
    print("P169:", p_liquidez(wb))
    print()
    print("P171:", p_razon_acida(wb))
    print()
    print("P173:", p_razon_endeudamiento(wb))
    print()
    print("P175:", p_estructura_deuda(wb))
    print()
    print("P179:", p_cobertura_gf(wb))
    print()
    print("P181:", p_rentabilidad_patrim_controladora(wb))
    print()
    print("P183:", p_rentabilidad_patrim_total(wb))
    print()
    print("P193:", p_activos_totales_intro(wb))
    print()
    print("P195:", p_activos_nocorrientes_bridge(wb))
    print()
    print("P197:", p_activos_corrientes_bridge(wb))
    print()
    print("P199:", p_pasivos_bridge(wb))
    print()
    print("P201:", p_patrimonio_bridge(wb))
