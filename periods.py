# -*- coding: utf-8 -*-
"""
Deriva las etiquetas de período (Jun26, Jun25, Dic25, fechas largas, temporadas, cantidad de meses del
tramo calendario y de la temporada) a partir de las fechas del Excel, para que el motor de generación
sirva para cualquiera de los 4 cierres trimestrales (marzo/trimestral, junio/semestral, septiembre/9M,
diciembre/anual) sin quedar amarrado a un trimestre fijo.
"""
import datetime

MESES_ABR = ['', 'Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']
MESES_LARGO = ['', 'enero','febrero','marzo','abril','mayo','junio','julio','agosto',
               'septiembre','octubre','noviembre','diciembre']
MESES_NUM_PALABRA = {3: 'tres', 6: 'seis', 9: 'nueve', 12: 'doce'}

def _short(d):
    return f"{MESES_ABR[d.month]}{d.year % 100:02d}"

def _long(d):
    return f"{d.day} de {MESES_LARGO[d.month]} de {d.year}"

def _season_labels(d):
    # Temporada agrícola: 1 jul - 30 jun. Si month<=6, la temporada que termina es (year-1)/(year).
    if d.month <= 6:
        y0, y1 = d.year - 1, d.year
    else:
        y0, y1 = d.year, d.year + 1
    cur = f"T{y0%100:02d}/{y1%100:02d}"
    prior = f"T{(y0-1)%100:02d}/{(y1-1)%100:02d}"
    return cur, prior, y0, y1

def derive_periods(wb, cur_date=None, prior_fy_date=None):
    """cur_date / prior_fy_date: datetime.date del cierre actual y del cierre de año fiscal anterior.
    Si no se entregan, se intentan leer de la hoja 'DF Neta' (C5 y E5)."""
    if cur_date is None or prior_fy_date is None:
        ws = wb['DF Neta']
        cur_date = ws['C5'].value
        prior_fy_date = ws['E5'].value
        if hasattr(cur_date, 'date'): cur_date = cur_date.date()
        if hasattr(prior_fy_date, 'date'): prior_fy_date = prior_fy_date.date()

    prior_same_period = cur_date.replace(year=cur_date.year - 1)
    season_cur, season_prior, season_y0, season_y1 = _season_labels(cur_date)

    # Tramo calendario (año en curso, Ene-cierre): marzo=3, junio=6, septiembre=9, diciembre=12.
    cal_meses = cur_date.month
    # Tramo temporada (temporada agrícola, Jul-cierre): junio=12 (temporada completa), resto = lo que
    # corresponda según cuántos meses van corridos desde el 1 de julio.
    season_meses = (cal_meses - 6) if cal_meses > 6 else (cal_meses + 6)

    # El párrafo preámbulo ("comparándose con los estados financieros al...") usa como referencia el
    # cierre anual anterior (Dic) para los cierres de semestre/año (junio y diciembre), y el mismo
    # período del año anterior para los cierres trimestrales intermedios (marzo y septiembre) — así
    # vienen redactados los informes originales de Hortifrut para cada tipo de cierre.
    if cal_meses in (6, 12):
        preambulo_ref_short = _short(prior_fy_date)
        preambulo_ref_long = _long(prior_fy_date)
    else:
        preambulo_ref_short = _short(prior_same_period)
        preambulo_ref_long = _long(prior_same_period)

    return {
        'CUR_SHORT': _short(cur_date),               # Jun26
        'PRIOR_SHORT': _short(prior_same_period),     # Jun25
        'FY_PRIOR_SHORT': _short(prior_fy_date),      # Dic25
        'CUR_LONG': _long(cur_date),                  # 30 de junio de 2026
        'FY_PRIOR_LONG': _long(prior_fy_date),         # 31 de diciembre de 2025
        'SEASON_CUR': season_cur,                     # T25/26
        'SEASON_PRIOR': season_prior,                 # T24/25
        'CAL_MESES': cal_meses,                       # 3 / 6 / 9 / 12
        'SEASON_MESES': season_meses,                 # 9 / 12 / 3 / 6 (según el cierre)
        'SEASON_MESES_PALABRA': MESES_NUM_PALABRA.get(season_meses, str(season_meses)),
        'CUR_MONTH_LOWER_YEAR': f"{MESES_LARGO[cur_date.month]} {cur_date.year}",      # "junio 2026"
        'CUR_MONTH_CAP': MESES_LARGO[cur_date.month].capitalize(),                    # "Junio"
        'SEASON_RANGE_LONG': f"julio {season_y0} – {MESES_LARGO[cur_date.month]} {cur_date.year}",
        'PREAMBULO_REF_SHORT': preambulo_ref_short,
        'PREAMBULO_REF_LONG': preambulo_ref_long,
    }

# Tokens literales usados dentro de los templates de ar_engine.py (siempre redactados en "estilo
# Jun26", que es el estilo de referencia pedido para todos los trimestres), mapeados a las claves de
# arriba. localize() los reemplaza por las etiquetas derivadas del Excel cargado, independiente de qué
# trimestre sea en realidad.
TOKEN_MAP = {
    '30 de junio de 2026': 'CUR_LONG',
    '31 de diciembre de 2025': 'FY_PRIOR_LONG',
    '12M T25/26': None,   # manejado abajo (compuesto, con el conteo de meses dinámico)
    '12M T24/25': None,
    'T25/26': 'SEASON_CUR',
    'T24/25': 'SEASON_PRIOR',
    'Jun26': 'CUR_SHORT',
    'Jun25': 'PRIOR_SHORT',
    'Dic25': 'FY_PRIOR_SHORT',
}

def localize(text, periods):
    # Importante: se reemplaza por marcadores temporales únicos antes de insertar el texto final, en
    # dos pasadas. Si se reemplazara directamente en una sola pasada, el resultado de un reemplazo
    # podría quedar "recapturado" por un reemplazo posterior cuyo patrón de búsqueda coincide con el
    # texto recién insertado — por ejemplo, en el cierre de diciembre CUR_SHORT es literalmente
    # "Dic25", el mismo texto que el token fijo 'Dic25' (que representa FY_PRIOR_SHORT) buscado más
    # abajo, lo que rompería la fecha recién puesta.
    repl = [
        ('12M T25/26', f"{periods['SEASON_MESES']}M {periods['SEASON_CUR']}"),
        ('12M T24/25', f"{periods['SEASON_MESES']}M {periods['SEASON_PRIOR']}"),
        # Tramo calendario: "6 meses" (único conteo usado en los templates de ar_engine.py) -> el
        # conteo real de este cierre (3/6/9/12).
        ('6 meses', f"{periods['CAL_MESES']} meses"),
        ('30 de junio de 2026', periods['CUR_LONG']),
        ('31 de diciembre de 2025', periods['FY_PRIOR_LONG']),
        ('T25/26', periods['SEASON_CUR']),
        ('T24/25', periods['SEASON_PRIOR']),
        ('Jun26', periods['CUR_SHORT']),
        ('Jun25', periods['PRIOR_SHORT']),
        ('Dic25', periods['FY_PRIOR_SHORT']),
    ]
    placeholders = []
    for i, (old, new) in enumerate(repl):
        ph = f"\x00{i}\x00"
        if old in text:
            text = text.replace(old, ph)
            placeholders.append((ph, new))
    for ph, new in placeholders:
        text = text.replace(ph, new)
    return text

_MES_LARGO_TO_NUM = {m: i for i, m in enumerate(MESES_LARGO) if m}

def parse_long_date(text):
    """Extrae una fecha tipo 'Al 31 de marzo de 2026' o '...al 31 de diciembre de 2025...' del texto,
    o None si no calza el patrón. Case-insensitive en el nombre del mes."""
    import re
    m = re.search(r'(\d{1,2}) de (\w+) de (\d{4})', text, re.IGNORECASE)
    if not m:
        return None
    day = int(m.group(1))
    mes_num = _MES_LARGO_TO_NUM.get(m.group(2).lower())
    if mes_num is None:
        return None
    year = int(m.group(3))
    try:
        return datetime.date(year, mes_num, day)
    except ValueError:
        return None

def derive_old_periods_from_title(title_text):
    """A partir del texto del título del documento BASE ('Al dd de mes de yyyy', ya en estilo Jun26
    antes de ser sobrescrito), reconstruye el período de ESE documento (el trimestre anterior), para
    poder barrer después cualquier mención suelta de sus fechas/etiquetas que haya quedado sin
    regenerar (secciones no cubiertas por un template, frases de contexto libres, etc.) y reemplazarla
    por la del trimestre nuevo. El cierre de año fiscal anterior de cualquier cierre es siempre el 31 de
    diciembre del año calendario anterior (confirmado contra los 4 cierres reales: Mar/Jun/Sep/Dic)."""
    old_cur_date = parse_long_date(title_text)
    if old_cur_date is None:
        return None
    old_prior_fy_date = datetime.date(old_cur_date.year - 1, 12, 31)
    return derive_periods(None, cur_date=old_cur_date, prior_fy_date=old_prior_fy_date)

def sweep_pairs(old_periods, new_periods):
    """Pares (texto viejo -> texto nuevo) para barrer, en párrafos que NO se regeneran por un template
    (secciones todavía no cubiertas, o frases de contexto libre cuya redacción es demasiado variable
    para anclar), cualquier mención suelta de fechas/etiquetas del trimestre del documento BASE que haya
    quedado sin actualizar. Si un valor viejo coincide con el nuevo (p.ej. FY_PRIOR_SHORT suele
    mantenerse igual de un trimestre a otro dentro del mismo año fiscal), el reemplazo es un no-op
    inofensivo. Se ordenan los compuestos ('NM T.../..') antes que sus partes sueltas."""
    pairs = [
        (f"{old_periods['SEASON_MESES']}M {old_periods['SEASON_CUR']}", f"{new_periods['SEASON_MESES']}M {new_periods['SEASON_CUR']}"),
        (f"{old_periods['SEASON_MESES']}M{old_periods['SEASON_CUR']}", f"{new_periods['SEASON_MESES']}M{new_periods['SEASON_CUR']}"),
        (f"{old_periods['SEASON_MESES']}M {old_periods['SEASON_PRIOR']}", f"{new_periods['SEASON_MESES']}M {new_periods['SEASON_PRIOR']}"),
        (f"{old_periods['SEASON_MESES']}M{old_periods['SEASON_PRIOR']}", f"{new_periods['SEASON_MESES']}M{new_periods['SEASON_PRIOR']}"),
        (old_periods['CUR_LONG'], new_periods['CUR_LONG']),
        (old_periods['FY_PRIOR_LONG'], new_periods['FY_PRIOR_LONG']),
        (old_periods['SEASON_CUR'], new_periods['SEASON_CUR']),
        (old_periods['SEASON_PRIOR'], new_periods['SEASON_PRIOR']),
        (old_periods['CUR_SHORT'], new_periods['CUR_SHORT']),
        (old_periods['PRIOR_SHORT'], new_periods['PRIOR_SHORT']),
        (old_periods['FY_PRIOR_SHORT'], new_periods['FY_PRIOR_SHORT']),
    ]
    # Sin duplicados ni pares no-op (viejo == nuevo ya es seguro de por sí, pero no vale la pena
    # procesarlos), preservando el orden (compuestos primero).
    seen = set()
    out = []
    for old, new in pairs:
        if old == new or old in seen:
            continue
        seen.add(old)
        out.append((old, new))
    return out

def sweep_text(text, old_periods, new_periods):
    """Reemplaza en `text` cualquier mención de las etiquetas de período del documento BASE (viejas)
    por las del trimestre nuevo, con el mismo esquema de dos pasadas por marcador que localize() (ver
    su comentario) para evitar recapturas entre reemplazos."""
    pairs = sweep_pairs(old_periods, new_periods)
    placeholders = []
    for i, (old, new) in enumerate(pairs):
        ph = f"\x01{i}\x01"
        if old in text:
            text = text.replace(old, ph)
            placeholders.append((ph, new))
    for ph, new in placeholders:
        text = text.replace(ph, new)
    return text

if __name__ == '__main__':
    import sys
    sys.path.insert(0, '.')
    import openpyxl
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    p = derive_periods(wb)
    for k, v in p.items():
        print(k, '=', v)
    print(localize("El EBITDA a Jun26 alcanzó X, comparado con Jun25 y Dic25. 12M T25/26 vs 12M T24/25. "
                    "Durante los 6 meses terminados en Jun26...", p))
